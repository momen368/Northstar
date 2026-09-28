from io import BytesIO

from docx import Document
from pypdf import PdfWriter
import pytest

from backend.app.services import resume_parser
from backend.app.services.resume_parser import (
    CorruptedDocumentError,
    EmptyDocumentError,
    UnsupportedResumeFormatError,
    extract_text,
    normalize_whitespace,
)
from backend.tests.pdf_fixture import text_pdf_bytes


def test_extracts_text_from_pdf(tmp_path):
    pdf_path = tmp_path / "resume.pdf"
    pdf_path.write_bytes(text_pdf_bytes("Professional Experience"))

    assert extract_text(pdf_path, "pdf") == "Professional Experience"


def test_extracts_docx_paragraphs_and_tables(tmp_path):
    document = Document()
    document.add_heading("Technical Skills", level=1)
    document.add_paragraph("Python and data engineering")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "SQL"
    table.cell(0, 1).text = "Cloud systems"
    docx_path = tmp_path / "resume.docx"
    document.save(docx_path)

    extracted = extract_text(docx_path, "docx")

    assert "Technical Skills" in extracted
    assert "Python and data engineering" in extracted
    assert "SQL | Cloud systems" in extracted


def test_empty_pdf_and_docx_raise_clear_errors(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    pdf_path = tmp_path / "empty.pdf"
    writer.write(pdf_path)
    empty_docx = Document()
    docx_path = tmp_path / "empty.docx"
    empty_docx.save(docx_path)

    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        extract_text(pdf_path, "pdf")
    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        extract_text(docx_path, "docx")


def test_corrupted_document_raises_clear_error(tmp_path):
    path = tmp_path / "corrupted.pdf"
    path.write_bytes(b"%PDF-1.4\nnot a valid PDF")

    with pytest.raises(CorruptedDocumentError, match="PDF resume could not be read"):
        extract_text(path, "pdf")


def test_parser_wraps_reader_errors(tmp_path, monkeypatch):
    path = tmp_path / "resume.pdf"
    path.write_bytes(text_pdf_bytes())

    def fail_reader(*_args, **_kwargs):
        raise OSError("low-level parser details")

    monkeypatch.setattr(resume_parser, "PdfReader", fail_reader)
    with pytest.raises(CorruptedDocumentError, match="PDF resume could not be read"):
        extract_text(path, "pdf")


def test_unsupported_file_type_has_clear_error(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("resume")

    with pytest.raises(UnsupportedResumeFormatError, match="only PDF and DOCX"):
        extract_text(path, "txt")


def test_normalizes_whitespace_without_collapsing_sections():
    extracted = normalize_whitespace("Summary\t text\n\n\nSkills\u00a0and tools  \n\n")

    assert extracted == "Summary text\n\nSkills and tools"