from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from backend.app.core.config import settings


class DocumentLoadError(ValueError):
    pass


@dataclass(frozen=True)
class LoadedDocument:
    source: str
    title: str
    category: str
    content: str


def clean_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    output = []
    for line in lines:
        if line:
            output.append(line)
        elif output and output[-1] != "":
            output.append("")
    return "\n".join(output).strip()


def load_document(file_path: str | Path, category: str | None = None) -> LoadedDocument:
    path = Path(file_path)
    if not path.is_file():
        raise DocumentLoadError("Knowledge document does not exist.")
    extension = path.suffix.casefold()
    try:
        if extension in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(str(path), strict=True)
            if reader.is_encrypted:
                raise DocumentLoadError("Encrypted knowledge PDFs are not supported.")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif extension == ".docx":
            with ZipFile(path) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise DocumentLoadError("DOCX document content is missing.")
            document = Document(path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            raise DocumentLoadError("Supported knowledge documents are TXT, Markdown, PDF, and DOCX.")
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError("Knowledge document is invalid or unreadable.") from exc
    content = clean_text(text)
    if not content:
        raise DocumentLoadError("Knowledge document contains no readable text.")
    inferred_category = category or (path.parent.name if path.parent.name else "general")
    resolved_path = path.resolve()
    try:
        source = resolved_path.relative_to(settings.knowledge_base_dir).as_posix()
    except ValueError:
        source = f"{path.parent.name}/{path.name}" if path.parent.name else path.name
    return LoadedDocument(source, path.stem, inferred_category, content)